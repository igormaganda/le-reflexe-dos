from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\Administrator\Downloads\Projets\reflexe-dos")
OUT = ROOT / "production"
IMG = ROOT / "assets" / "images"

BLACK = RGBColor(24, 29, 28)
GREEN = RGBColor(13, 59, 51)
CORAL = RGBColor(220, 103, 75)
GREY = RGBColor(92, 99, 97)
LIGHT = "F2F4F1"
WHITE = RGBColor(255, 255, 255)

HAS_URL = "https://www.has-sante.fr/jcms/c_2961499/"
AMELI_MOVE_URL = "https://www.ameli.fr/assure/sante/themes/lombalgie-aigue/traitement-prevention"
AMELI_INFO_URL = "https://www.ameli.fr/assure/sante/themes/lombalgie-aigue/comprendre-lombalgie"
AMELI_CHRONIC_URL = "https://www.ameli.fr/assure/sante/themes/lombalgie-aigue/lombalgie-subaigue-chronique"
WHO_URL = "https://www.who.int/publications/i/item/9789240081789"


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text, url, color="0D3B33", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def base_document(short_title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.65)
    sec.bottom_margin = Cm(1.55)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, before, after in (("Title", 31, 0, 11), ("Heading 1", 22, 0, 10), ("Heading 2", 15, 12, 5), ("Heading 3", 11, 9, 3)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = sec.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"LE RÉFLEXE DOS   ·   {short_title.upper()}")
    set_run_font(r, size=7.5, bold=True, color=GREY)

    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Document de travail 0.1   •   Relecture médicale obligatoire   •   ")
    set_run_font(r, size=7.5, color=GREY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)

    core = doc.core_properties
    core.author = "Le Réflexe Dos"
    core.subject = "Prototype éditorial destiné à une relecture médicale et utilisateur"
    core.comments = "Version 0.1 produite le 4 octobre 2026"
    return doc


def add_kicker(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    set_run_font(r, size=8.5, bold=True, color=CORAL)
    return p


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(8)
    p.add_run(text)
    if subtitle:
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(12)
        r = q.add_run(subtitle)
        set_run_font(r, name="Aptos Display", size=14, color=GREY)


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
    p.add_run(text)
    return p


def add_bullets(doc, items, checkbox=False):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        prefix = "☐ " if checkbox else ""
        p.add_run(prefix + item)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.28)
        p.add_run(f"{index}. {item}")


def add_note(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.right_indent = Cm(0.25)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "DC674B")
    borders.append(left)
    pPr.append(borders)
    r = p.add_run(label.upper() + "  ")
    set_run_font(r, size=8.5, bold=True, color=CORAL)
    p.add_run(text)


def add_lines(doc, prompts, line_count=1):
    for prompt in prompts:
        p = doc.add_paragraph()
        r = p.add_run(prompt)
        set_run_font(r, bold=True, size=9.5)
        for _ in range(line_count):
            q = doc.add_paragraph("________________________________________________________________________________")
            q.paragraph_format.space_after = Pt(4)
            for run in q.runs:
                set_run_font(run, size=8, color=RGBColor(180, 184, 182))


def add_image(doc, filename, caption, width=17.0, height=None):
    path = IMG / filename
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        if height is not None:
            run.add_picture(str(path), height=Cm(height))
        else:
            run.add_picture(str(path), width=Cm(width))
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        for r in cap.runs:
            set_run_font(r, size=8, color=GREY)


def add_sources(doc):
    add_heading(doc, "Sources de référence", 1)
    sources = [
        ("Haute Autorité de santé  Prise en charge du patient présentant une lombalgie commune", HAS_URL),
        ("Assurance Maladie  Lombalgie ou mal de dos", AMELI_INFO_URL),
        ("Assurance Maladie  Le bon traitement est le mouvement", AMELI_MOVE_URL),
        ("Assurance Maladie  Éviter le passage à la chronicité", AMELI_CHRONIC_URL),
        ("Organisation mondiale de la Santé  Guideline on chronic primary low back pain", WHO_URL),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, label, url)
    add_body(doc, "Sources consultées le 4 octobre 2026. Le manuscrit reformule les informations. Il ne reproduit pas les recommandations et ne se substitue pas à leur lecture par le relecteur médical.")


def new_page(doc, kicker, title, subtitle=None):
    if len(doc.paragraphs) > 0:
        doc.add_page_break()
    add_kicker(doc, kicker)
    add_title(doc, title, subtitle)


def cover(doc, kicker, title, subtitle, image, caption, audience, version="Version de travail 0.1"):
    add_kicker(doc, kicker)
    add_title(doc, title, subtitle)
    cover_height = 14.2 if image == "couverture-praticien-patiente.png" else None
    add_image(doc, image, caption, width=17.0, height=cover_height)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run(version.upper())
    set_run_font(r, size=8.5, bold=True, color=CORAL)
    add_body(doc, audience)
    add_note(doc, "Statut", "Manuscrit de production. Les messages de santé, exercices et illustrations anatomiques doivent être relus par un médecin ou un kinésithérapeute avant diffusion.")


def add_safety_page(doc):
    new_page(doc, "Sécurité", "Avant de poursuivre", "Ce document informe et accompagne mais ne pose aucun diagnostic")
    add_body(doc, "Votre douleur mérite d’être prise au sérieux. Elle est souvent liée à une lombalgie commune, mais un guide ne peut pas vérifier votre situation. Si la douleur est récente, inhabituelle ou change nettement, demandez un avis médical. Vous n’avez pas à prendre cette décision seul.")
    add_heading(doc, "Demandez une aide urgente", 2)
    add_bullets(doc, [
        "difficulté nouvelle à uriner, fuites urinaires ou perte du contrôle des selles",
        "perte de sensibilité autour du périnée ou sensation d’anesthésie en selle",
        "faiblesse marquée ou qui progresse dans une ou deux jambes",
        "douleur apparue après un traumatisme important",
        "douleur accompagnée d’une altération importante de l’état général, d’une fièvre ou de symptômes respiratoires inhabituels",
    ])
    add_body(doc, "En France, appelez le 15 ou le 112 si la situation paraît urgente ou si vous ne savez pas vers quel service vous tourner.")
    add_heading(doc, "Prenez rapidement un avis médical", 2)
    add_bullets(doc, [
        "douleur qui s’aggrave régulièrement, persiste au repos ou réveille sans cesse la nuit",
        "perte de poids inexpliquée, antécédent de cancer, traitement prolongé par corticoïdes ou immunodépression",
        "engourdissement ou faiblesse qui apparaît ou augmente dans une jambe",
        "douleur qui ne s’améliore pas ou qui limite fortement les activités malgré quelques jours d’adaptation",
    ])
    add_note(doc, "Repère", "Un signe isolé ne veut pas forcément dire qu’il existe une maladie grave. Il mérite simplement d’être évalué dans son contexte. En cas de doute, un professionnel de santé peut vous aider à choisir la bonne suite.")


def build_ebook():
    doc = base_document("Comprendre son dos")
    cover(
        doc,
        "Ebook pédagogique",
        "Comprendre son dos",
        "Des repères simples pour sortir du flou sans dramatiser",
        "couverture-praticien-patiente.png",
        "Un praticien explique la région lombaire à une patiente  Visuel de couverture validé pour la phase de production",
        "Pour les adultes qui veulent comprendre une lombalgie commune, préparer une consultation et reprendre progressivement leurs activités.",
    )
    add_safety_page(doc)

    new_page(doc, "Mode d’emploi", "Comment lire ce guide", "Chaque chapitre répond à une décision concrète")
    add_numbered(doc, [
        "Comprendre ce qui est fréquent et ce qui mérite un avis médical.",
        "Vérifier les signaux d’alerte avant d’essayer une stratégie d’autogestion.",
        "Choisir une activité utile dans votre journée réelle.",
        "Observer la réponse du corps sans transformer chaque sensation en test.",
        "Ajuster la dose et demander de l’aide si la situation sort du cadre prévu.",
    ])
    add_heading(doc, "Ce que vous pouvez attendre", 2)
    add_body(doc, "Le guide vous donne des mots simples, une méthode de décision et des exemples. Il ne promet pas de supprimer la douleur en quelques jours. Il vous aide à avancer à votre rythme et à reprendre une part d’activité adaptée à votre situation.")
    add_note(doc, "Notre engagement", "Nous croyons votre douleur. Nous ne vous demandons ni de la minimiser ni de forcer. Nous vous aidons à faire un prochain pas raisonnable, puis à ajuster selon votre réponse.")
    add_heading(doc, "Ce que le guide ne fait pas", 2)
    add_bullets(doc, [
        "diagnostiquer la cause d’une douleur",
        "interpréter une imagerie ou modifier un traitement",
        "remplacer un examen clinique",
        "prescrire un exercice personnalisé",
    ])
    add_lines(doc, ["Ce que je veux pouvoir refaire plus facilement", "La question que je veux poser à un professionnel"], 2)

    new_page(doc, "Sommaire", "Le parcours éditorial")
    sections = [
        ("01", "La douleur et les dommages", "Comprendre pourquoi l’intensité d’une douleur ne mesure pas à elle seule l’état des tissus"),
        ("02", "Un dos fait pour bouger", "Situer le rôle des vertèbres, muscles, nerfs et habitudes de vie"),
        ("03", "La sécurité avant les exercices", "Reconnaître les situations qui sortent de l’autogestion"),
        ("04", "Reprendre une activité", "Choisir une dose réaliste et l’ajuster"),
        ("05", "Sommeil stress et attention", "Comprendre ce qui peut amplifier ou apaiser l’expérience"),
        ("06", "Consultation imagerie et questions", "Préparer un échange utile avec un professionnel"),
        ("07", "Prévenir les récidives", "Construire une continuité sans dépendre d’une routine parfaite"),
    ]
    for num, title, desc in sections:
        p = doc.add_paragraph()
        r = p.add_run(num + "  ")
        set_run_font(r, size=9, bold=True, color=CORAL)
        r = p.add_run(title)
        set_run_font(r, size=12, bold=True)
        q = doc.add_paragraph(desc)
        q.paragraph_format.left_indent = Cm(0.85)
        q.paragraph_format.space_after = Pt(7)

    new_page(doc, "Chapitre 01", "La douleur et les dommages", "Une douleur forte est réelle sans être une mesure exacte de la lésion")
    add_body(doc, "La douleur est une expérience de protection. Elle mobilise le système nerveux lorsqu’il estime qu’une partie du corps est menacée. Cette estimation utilise des informations venant des tissus, mais aussi le contexte, les expériences passées, la fatigue, le stress et le sens donné à la situation.")
    add_body(doc, "Cela ne rend pas la douleur imaginaire. Cela explique pourquoi deux journées comparables peuvent être ressenties différemment et pourquoi une même image radiologique ne produit pas les mêmes symptômes chez tout le monde.")
    add_heading(doc, "Ce que la douleur permet de savoir", 2)
    add_bullets(doc, [
        "quelque chose mérite votre attention maintenant",
        "certaines activités sont devenues difficiles ou inquiétantes",
        "votre stratégie doit tenir compte du sommeil, du travail et de la confiance, pas seulement d’un chiffre de douleur",
    ])
    add_heading(doc, "Ce qu’elle ne permet pas de conclure seule", 2)
    add_bullets(doc, [
        "la quantité exacte de dommages dans le dos",
        "la nécessité automatique d’une imagerie",
        "l’existence d’un mouvement interdit pour toujours",
        "le délai précis de récupération",
    ])
    add_note(doc, "Repère", "La Haute Autorité de santé rappelle l’absence de corrélation systématique entre les symptômes et les signes observés à l’imagerie. L’examen clinique et l’évolution des symptômes restent essentiels.")

    new_page(doc, "Chapitre 01", "Observer sans surveiller chaque sensation", "Suivre la fonction donne souvent plus d’informations qu’un score isolé")
    add_body(doc, "Noter uniquement la douleur peut donner l’impression que la journée entière doit être organisée autour d’elle. Ajoutez des indicateurs de capacité. Ils montrent ce que vous récupérez, même lorsque la douleur varie encore.")
    add_heading(doc, "Quatre indicateurs utiles", 2)
    add_bullets(doc, [
        "une activité importante réalisée aujourd’hui",
        "le temps passé dans une même position avant d’avoir besoin de changer",
        "la confiance ressentie avant et après une activité",
        "la récupération observée le soir et le lendemain",
    ])
    add_heading(doc, "Exercice de mise au point", 2)
    add_lines(doc, [
        "Une activité que j’évite actuellement",
        "Ce que je crains qu’il arrive",
        "Une version plus courte ou plus simple que je pourrais discuter ou tester",
        "Le signe qui me dira de réduire ou de demander conseil",
    ], 1)
    add_note(doc, "À retenir", "Une progression n’est pas forcément linéaire. Une journée plus douloureuse n’efface pas les capacités gagnées. Elle invite à regarder la dose, le contexte et la récupération.")

    new_page(doc, "Chapitre 02", "Un dos fait pour bouger", "La colonne combine stabilité adaptation et mouvement")
    add_image(doc, "exercice-mobilite.png", "Mouvement guidé et progressif  L’illustration finale devra montrer une amplitude confortable et des adaptations visibles", width=14.5)
    add_body(doc, "La région lombaire comprend cinq vertèbres, des disques, de petites articulations, des ligaments, des muscles et des nerfs. L’ensemble supporte des charges variées et s’adapte aux gestes répétés. Le dos n’est pas une pile de pièces qu’un mouvement ordinaire remettrait facilement de travers.")
    add_heading(doc, "Charge et capacité", 2)
    add_body(doc, "Une charge n’est pas seulement un poids. Elle inclut la durée, la répétition, la vitesse, la position, la fatigue et la confiance. La capacité correspond à ce que vous tolérez aujourd’hui. Elle peut évoluer grâce à une exposition progressive.")
    add_bullets(doc, [
        "Si la dose est très inférieure à votre capacité, le geste paraît facile.",
        "Si elle s’en approche, vous pouvez ressentir un effort ou une appréhension gérable.",
        "Si elle la dépasse nettement, la réponse peut durer et nécessiter une réduction temporaire.",
    ])

    new_page(doc, "Chapitre 03", "La sécurité avant les exercices", "Décider si l’autogestion est adaptée aujourd’hui")
    add_body(doc, "Commencez chaque nouveau programme par la page de sécurité. La plupart des lombalgies sont communes, mais certains signes imposent un examen plutôt qu’une expérimentation à domicile.")
    add_heading(doc, "Le contrôle en trois questions", 2)
    add_numbered(doc, [
        "Un signe neurologique nouveau est il apparu, comme une faiblesse qui progresse, une difficulté à uriner ou une perte de sensibilité du périnée",
        "La douleur est elle liée à un traumatisme important, à une fièvre, à une altération générale ou à un contexte médical particulier",
        "Les symptômes ont ils changé de nature ou se sont ils aggravés de façon inhabituelle",
    ])
    add_body(doc, "Si la réponse est oui ou incertaine, suspendez le programme et demandez un avis adapté. Si la réponse est non, choisissez une action légère et observable plutôt qu’un test maximal.")
    add_heading(doc, "La règle d’arrêt pendant une activité", 2)
    add_bullets(doc, [
        "arrêtez immédiatement en cas de faiblesse nouvelle, malaise, douleur thoracique, essoufflement inhabituel ou perte de sensibilité",
        "réduisez l’amplitude, la durée ou la répétition si la réponse devient nettement plus forte ou moins contrôlable",
        "demandez conseil si l’aggravation persiste ou modifie vos activités de façon importante",
    ])

    new_page(doc, "Chapitre 04", "Reprendre une activité", "Choisir une dose que vous pouvez répéter")
    add_body(doc, "Les recommandations françaises encouragent la poursuite des activités quotidiennes autant que possible et une activité physique adaptée. Vous n’avez rien à prouver aujourd’hui. L’idée est de retrouver peu à peu une continuité qui vous convient.")
    add_heading(doc, "La plus petite version utile", 2)
    add_numbered(doc, [
        "Choisissez une activité qui compte dans votre vie, par exemple marcher jusqu’à la boulangerie, vous habiller seul ou rester à votre poste pendant une réunion.",
        "Réduisez un paramètre, comme la durée, la distance, la charge ou le nombre de répétitions.",
        "Gardez une marge pour pouvoir refaire l’action demain ou après demain.",
        "Observez la réponse pendant l’activité, le soir et le lendemain.",
        "Augmentez un seul paramètre lorsque la réponse est acceptable et prévisible.",
    ])
    add_heading(doc, "Exemple", 2)
    add_body(doc, "Si dix minutes de marche semblent incertaines, commencez par deux sorties de quatre minutes à allure habituelle. Après deux ou trois essais, augmentez soit la durée d’une sortie, soit l’allure. Ne changez pas tout à la fois.")
    add_note(doc, "Important", "Cet exemple illustre une méthode de dosage. Il ne constitue pas une prescription. Une douleur irradiée, une faiblesse, un contexte médical complexe ou une reprise difficile justifient un conseil personnalisé.")

    new_page(doc, "Chapitre 05", "Sommeil stress et attention", "Le contexte peut modifier l’intensité et la récupération")
    add_body(doc, "Un sommeil insuffisant, une forte inquiétude ou une journée sans possibilité de changer de position peuvent augmenter la sensibilité. Ces facteurs n’expliquent pas tout et ne signifient pas que la douleur est psychologique. Ils font partie des informations utilisées par le système de protection.")
    add_heading(doc, "Trois leviers sobres", 2)
    add_bullets(doc, [
        "préparer une position de repos confortable sans chercher une posture parfaite",
        "alterner les positions et fractionner les tâches longues",
        "définir une action concrète pour le lendemain afin de réduire l’incertitude",
    ])
    add_heading(doc, "Question de fin de journée", 2)
    add_lines(doc, ["Ce qui a rendu la journée plus facile", "Ce qui a augmenté la difficulté", "Le réglage que je teste demain"], 1)

    new_page(doc, "Chapitre 06", "Préparer une consultation", "Arriver avec des informations utiles plutôt qu’un récit parfait")
    add_image(doc, "consultation-kine.png", "Consultation et décision partagée  Le professionnel recherche les signes d’alerte puis adapte les conseils", width=14.2)
    add_heading(doc, "Les informations à noter", 2)
    add_bullets(doc, [
        "date et circonstances d’apparition ou de modification",
        "zone de douleur et éventuelle irradiation dans une jambe",
        "faiblesse, engourdissement, troubles urinaires ou autres symptômes associés",
        "activités devenues difficiles et activités encore possibles",
        "médicaments, antécédents, examens déjà réalisés et attentes principales",
    ])
    add_heading(doc, "Questions possibles", 2)
    add_bullets(doc, [
        "Quels signes doivent me faire reconsulter rapidement",
        "Quelles activités puis je conserver dès maintenant",
        "Comment dois je ajuster la charge ou la durée",
        "Dans mon cas, qu’est ce qui justifierait une imagerie ou une orientation",
    ])

    new_page(doc, "Chapitre 07", "Prévenir les récidives", "Construire une continuité compatible avec les semaines imparfaites")
    add_body(doc, "La prévention repose moins sur un exercice secret que sur une activité régulière, choisie selon vos préférences et suffisamment souple pour survivre aux périodes chargées. Le programme doit prévoir une version courte, une version habituelle et une façon de reprendre après une interruption.")
    add_heading(doc, "Votre plan de continuité", 2)
    add_lines(doc, [
        "Mon activité principale",
        "Ma version de dix minutes ou moins",
        "Le moment réaliste dans ma semaine",
        "La personne ou le repère qui m’aide à reprendre",
        "Le signe qui m’indique de demander un avis professionnel",
    ], 1)
    add_note(doc, "Suite proposée", "Le cahier Dos solide 28 jours transforme ces repères en pratique quotidienne. Le guide SOS Lumbago 7 jours concerne les premiers jours d’une poussée et commence toujours par un contrôle de sécurité.")

    new_page(doc, "Références", "Sources et validation éditoriale")
    add_sources(doc)
    add_heading(doc, "Validation avant publication", 2)
    add_bullets(doc, [
        "relecture médicale de chaque message de sécurité et de chaque exemple",
        "validation des illustrations par un professionnel compétent",
        "test de compréhension auprès de lecteurs non spécialistes",
        "date de révision affichée dans le produit final",
    ])
    out = OUT / "01-Comprendre-son-dos-manuscrit-v0.1.docx"
    doc.save(out)
    return out


def build_workbook():
    doc = base_document("Dos solide 28 jours")
    cover(
        doc,
        "Cahier d’exercices",
        "Dos solide 28 jours",
        "Un parcours progressif pour retrouver de la régularité et de la confiance",
        "exercice-mobilite.png",
        "Mouvement simple réalisé dans une amplitude confortable  Démonstrations finales à valider cliniquement",
        "Pour les adultes dont la situation a été évaluée ou ne présente pas de signe d’alerte et qui souhaitent structurer une reprise progressive.",
    )
    add_safety_page(doc)

    new_page(doc, "Mode d’emploi", "Votre contrat pour 28 jours", "La régularité compte davantage que la performance")
    add_bullets(doc, [
        "Consacrez cinq à quinze minutes au rendez vous du jour.",
        "Choisissez le niveau le plus facile qui vous permet de rester régulier.",
        "Ne rattrapez pas une séance manquée en doublant la suivante.",
        "Notez une capacité et une réaction, pas seulement l’intensité de la douleur.",
        "Suspendez le programme et demandez conseil si un signe d’alerte apparaît.",
    ])
    add_note(doc, "À votre rythme", "Une journée manquée n’est pas un échec. Reprenez par la version courte dès que votre situation le permet. Le cahier doit vous soutenir, pas ajouter une obligation.")
    add_heading(doc, "Les trois versions", 2)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Version courte", "Version habituelle", "Version plus longue"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        shade_cell(cell, "0D3B33")
        for r in cell.paragraphs[0].runs:
            set_run_font(r, size=9, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    values = ["2 à 5 minutes\nPour les jours difficiles", "6 à 10 minutes\nPour la plupart des jours", "11 à 15 minutes\nSi la réponse reste prévisible"]
    row = table.add_row()
    prevent_row_split(row)
    for i, value in enumerate(values):
        row.cells[i].text = value
        set_cell_margins(row.cells[i])
        row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in row.cells[i].paragraphs[0].runs:
            set_run_font(r, size=9)
    add_note(doc, "Validation", "Les durées et exercices sont des choix éditoriaux de travail. Un professionnel de santé doit confirmer leur pertinence et proposer les adaptations avant publication.")

    new_page(doc, "Point de départ", "Votre situation aujourd’hui")
    add_lines(doc, [
        "L’activité importante que je veux retrouver",
        "Ce que je peux déjà faire même brièvement",
        "Ce que j’évite et la raison",
        "Le moment le plus réaliste pour mon rendez vous quotidien",
        "La personne que je contacterai si la situation change",
    ], 1)
    add_heading(doc, "Mes repères de départ", 2)
    add_bullets(doc, [
        "Confiance pour bouger aujourd’hui  ____ sur 10",
        "Temps toléré dans mon activité cible  ____________________",
        "Qualité du sommeil cette semaine  faible  moyenne  bonne",
        "Impact sur le travail ou les tâches courantes  faible  moyen  fort",
    ], checkbox=True)

    new_page(doc, "Vue d’ensemble", "Le programme en quatre semaines")
    weeks = [
        ("Semaine 1", "Retrouver des repères", "Observer, fractionner, reprendre un mouvement familier"),
        ("Semaine 2", "Augmenter la tolérance", "Faire progresser un seul paramètre et récupérer"),
        ("Semaine 3", "Renforcer la confiance", "Réintroduire des gestes évités avec une échelle graduée"),
        ("Semaine 4", "Construire la continuité", "Préparer les semaines chargées et les futures poussées"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(("Période", "Cap", "Travail principal")):
        cell = table.rows[0].cells[i]
        cell.text = h
        shade_cell(cell, "0D3B33")
        for r in cell.paragraphs[0].runs:
            set_run_font(r, bold=True, color=WHITE, size=9)
    set_repeat_table_header(table.rows[0])
    for week, cap, work in weeks:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate((week, cap, work)):
            row.cells[i].text = value
            set_cell_margins(row.cells[i])
            for r in row.cells[i].paragraphs[0].runs:
                set_run_font(r, size=9)
    add_heading(doc, "Le bilan hebdomadaire", 2)
    add_body(doc, "À la fin de chaque semaine, répondez à trois questions. Qu’est devenu plus facile. Qu’est resté imprévisible. Quel seul paramètre allez vous modifier. Cette boucle reprend le principe utile de l’Agenda RUNWEEK sans transformer le cahier en calendrier chargé.")

    new_page(doc, "Semaine 01", "Jour 1 choisir un point d’appui", "Commencer par une action déjà possible")
    add_body(doc, "Choisissez une activité familière que vous pouvez faire sans vous tester. Cela peut être marcher dans le logement, préparer une boisson debout ou effectuer un déplacement court. La version du jour doit vous laisser une marge.")
    add_heading(doc, "Séance prototype", 2)
    add_numbered(doc, [
        "Vérifiez la page de sécurité et votre état général.",
        "Choisissez une activité de deux à cinq minutes.",
        "Commencez plus lentement que nécessaire pendant trente secondes.",
        "Poursuivez si le mouvement reste contrôlable et arrêtez avant l’épuisement.",
        "Observez la réponse une heure plus tard et le lendemain.",
    ])
    add_lines(doc, ["Mon activité", "Ma durée", "Ce qui m’a rassuré", "Le réglage pour demain"], 1)

    new_page(doc, "Semaine 01", "Jour 2 interrompre l’immobilité", "Changer de position avant que la journée ne se fige")
    add_body(doc, "Une position n’est pas mauvaise en elle même. La difficulté vient souvent de la durée, de la répétition ou du manque de choix. Aujourd’hui, installez trois rappels de changement de position dans une période habituellement longue.")
    add_heading(doc, "Trois micro pauses", 2)
    add_bullets(doc, [
        "se lever et faire quelques pas",
        "s’appuyer différemment ou alterner assis et debout",
        "réaliser un mouvement connu dans une amplitude confortable",
    ], checkbox=True)
    add_body(doc, "Chaque pause peut durer moins de deux minutes. Elle sert à varier, pas à corriger une posture supposée parfaite.")
    add_lines(doc, ["Moment 1 et réponse", "Moment 2 et réponse", "Moment 3 et réponse", "La variation la plus utile"], 1)

    new_page(doc, "Semaine 01", "Jour 3 respirer et relâcher l’effort inutile", "Distinguer protection et crispation")
    add_body(doc, "La douleur peut vous amener à retenir votre respiration ou à contracter en permanence le ventre, les épaules et les mains. L’exercice du jour vise seulement à retrouver un souffle libre dans une position confortable.")
    add_heading(doc, "Séance prototype", 2)
    add_numbered(doc, [
        "Installez vous assis, debout ou allongé selon votre préférence.",
        "Laissez l’expiration durer naturellement sans chercher une grande inspiration.",
        "Relâchez les épaules et la mâchoire pendant cinq cycles calmes.",
        "Reprenez une activité courante pendant une à deux minutes en gardant un souffle fluide.",
    ])
    add_note(doc, "Arrêt", "Interrompez l’exercice en cas de vertige, malaise ou essoufflement inhabituel. Il ne s’agit pas d’une technique respiratoire médicale.")
    add_lines(doc, ["Position choisie", "Avant l’exercice", "Après l’exercice", "Ce que j’ai remarqué"], 1)

    new_page(doc, "Semaine 01", "Jour 4 se lever d’une chaise", "Réintroduire un geste quotidien avec plusieurs options")
    add_body(doc, "Choisissez une chaise stable. Avancez les pieds jusqu’à trouver un appui confortable. Vous pouvez utiliser les accoudoirs, les cuisses ou un meuble stable. Penchez légèrement le tronc si cela vous aide, puis poussez dans les pieds pour vous lever.")
    add_heading(doc, "Options de travail", 2)
    add_bullets(doc, [
        "Version courte  une à trois répétitions espacées",
        "Version habituelle  deux séries courtes avec une pause",
        "Adaptation  chaise plus haute ou appui des mains",
    ])
    add_body(doc, "Le but est de retrouver des options, pas d’imposer une technique unique. Ne poursuivez pas si une faiblesse nouvelle, un engourdissement marqué ou une douleur difficile à contrôler apparaît.")
    add_lines(doc, ["Option utilisée", "Nombre confortable", "Réponse une heure après", "Décision pour la prochaine fois"], 1)

    new_page(doc, "Semaine 01", "Jour 5 marcher à dose choisie", "Utiliser la marche comme activité graduée")
    add_body(doc, "Choisissez un trajet simple, avec une possibilité de retour ou de pause. Décidez de la durée avant de partir. Une petite boucle répétable apporte plus d’informations qu’une longue sortie faite pour se tester.")
    add_heading(doc, "Plan du jour", 2)
    add_bullets(doc, [
        "trajet sécurisé et connu",
        "durée prévue et point de demi tour",
        "allure qui permet de parler normalement",
        "solution de repli si la fatigue augmente",
    ], checkbox=True)
    add_lines(doc, ["Durée prévue", "Durée réalisée", "Confiance avant et après", "Réponse le lendemain"], 1)
    add_note(doc, "Ajustement", "Si la réponse reste acceptable et prévisible, conservez la dose ou augmentez un seul paramètre lors de la prochaine marche. Si elle persiste nettement, réduisez la durée ou demandez conseil.")

    new_page(doc, "Semaine 01", "Jour 6 construire une échelle de confiance", "Découper un geste évité en étapes")
    add_body(doc, "Choisissez un geste important mais non urgent. Classez cinq versions de la plus facile à la plus difficile. Vous n’avez pas à atteindre la dernière aujourd’hui.")
    add_lines(doc, [
        "Geste choisi",
        "Étape 1 très accessible",
        "Étape 2",
        "Étape 3",
        "Étape 4",
        "Étape 5 objectif futur",
    ], 1)
    add_heading(doc, "Essai du jour", 2)
    add_body(doc, "Testez seulement l’étape 1, ou l’étape 2 si la première est déjà familière. Notez ce qui s’est réellement passé, distinctement de ce que vous craigniez.")
    add_lines(doc, ["Ce que je prévoyais", "Ce qui s’est passé", "Ma prochaine étape"], 1)

    new_page(doc, "Semaine 01", "Jour 7 récupérer et faire le bilan", "Conserver ce qui mérite d’être répété")
    add_heading(doc, "Bilan de la semaine", 2)
    add_lines(doc, [
        "L’action la plus facile à répéter",
        "Une capacité qui a progressé",
        "Une situation encore imprévisible",
        "Le réglage qui a le mieux fonctionné",
        "La question à poser si j’ai besoin d’aide",
    ], 1)
    add_heading(doc, "Décision pour la semaine 2", 2)
    add_bullets(doc, [
        "Je conserve la même dose pour stabiliser.",
        "J’augmente un seul paramètre de manière modeste.",
        "Je réduis et je demande conseil car la réponse reste difficile à contrôler.",
    ], checkbox=True)

    for week, title, goal, days in [
        ("Semaine 02", "Augmenter la tolérance", "Faire progresser une durée, une distance ou une répétition sans changer tous les paramètres", ["reprendre la marche", "varier les positions", "porter un objet léger", "répéter le lever de chaise", "activité choisie", "échelle de confiance", "bilan"]),
        ("Semaine 03", "Renforcer la confiance", "Réintroduire progressivement deux gestes évités et préparer une activité réelle", ["choisir le geste", "étape accessible", "répétition calme", "contexte réel", "récupération", "nouvelle étape", "bilan"]),
        ("Semaine 04", "Construire la continuité", "Créer une routine minimale et un plan de reprise après interruption", ["version courte", "semaine chargée", "activité plaisir", "travail et déplacements", "signal de reprise", "plan de poussée", "bilan final"]),
    ]:
        new_page(doc, week, title, goal)
        add_body(doc, "Cette page fixe le parcours de la prochaine version. Les séances détaillées seront rédigées après relecture des sept prototypes de la première semaine et validation des règles de dosage.")
        add_heading(doc, "Les sept rendez vous", 2)
        add_numbered(doc, [d.capitalize() for d in days])
        add_heading(doc, "Critère de progression", 2)
        add_body(doc, "La progression sera basée sur la capacité à répéter l’activité et sur la récupération observée, pas sur l’obligation d’obtenir une journée sans douleur.")
        add_lines(doc, ["Cap de la semaine", "Activité suivie", "Un paramètre à faire évoluer", "Condition pour réduire ou demander conseil"], 1)

    new_page(doc, "Jour 28", "Bilan final et suite")
    add_lines(doc, [
        "Ce que je peux faire aujourd’hui par rapport au jour 1",
        "L’activité que je veux maintenir deux fois par semaine",
        "Ma version courte pour les semaines difficiles",
        "Mes premiers signes de surcharge",
        "Mon plan de reprise après une interruption",
        "Le professionnel que je consulterai si la situation change",
    ], 1)
    add_heading(doc, "Mes indicateurs", 2)
    add_bullets(doc, [
        "Confiance pour bouger aujourd’hui  ____ sur 10",
        "Temps toléré dans mon activité cible  ____________________",
        "Nombre de jours actifs cette semaine  ____ sur 7",
    ], checkbox=True)

    new_page(doc, "Références", "Sources et validation éditoriale")
    add_sources(doc)
    add_heading(doc, "Étapes de production restantes", 2)
    add_bullets(doc, [
        "validation clinique des séances et des règles d’ajustement",
        "illustrations des options de mouvement avec consignes visuelles",
        "test de sept jours avec un petit groupe de lecteurs",
        "rédaction détaillée des semaines 2 à 4 après retours",
    ])
    out = OUT / "02-Dos-solide-28-jours-manuscrit-v0.1.docx"
    doc.save(out)
    return out


def build_sos():
    doc = base_document("SOS Lumbago 7 jours")
    cover(
        doc,
        "Guide pratique illustré",
        "SOS Lumbago 7 jours",
        "Des décisions simples pour traverser les premiers jours et préparer la suite",
        "couverture-praticien-patiente.png",
        "Un praticien montre la région lombaire à une patiente  Couverture retenue pour le produit",
        "Pour un adulte confronté à une poussée récente de douleur lombaire et qui cherche un cadre prudent jusqu’à l’amélioration ou la consultation.",
    )
    add_safety_page(doc)

    new_page(doc, "Lecture rapide", "Que faire maintenant", "Les cinq premières décisions")
    add_numbered(doc, [
        "Vérifiez les signes d’alerte de la page précédente.",
        "Choisissez une position qui vous permet de respirer et de vous détendre quelques minutes.",
        "Évitez de rester au lit toute la journée sauf consigne médicale particulière.",
        "Conservez de petits déplacements et les activités tolérables.",
        "Notez ce qui a changé et demandez un avis si les symptômes sont inhabituels, s’aggravent ou limitent fortement la vie quotidienne.",
    ])
    add_heading(doc, "Votre seul objectif aujourd’hui", 2)
    add_body(doc, "Trouvez une combinaison de repos bref, de positions variées et de mouvements simples que vous pouvez répéter sans vous mettre à l’épreuve. Un petit mieux compte. Vous n’avez pas à résoudre toute la semaine en une journée.")
    add_note(doc, "Vous guider", "La douleur peut rendre chaque choix plus difficile. Suivez une étape à la fois. Si une étape ne vous convient pas, adaptez la ou demandez de l’aide.")
    add_lines(doc, ["L’action essentielle d’aujourd’hui", "La personne à prévenir si j’ai besoin d’aide"], 1)

    new_page(doc, "Jour 0", "Les trente premières minutes", "Réduire l’incertitude avant de chercher un exercice")
    add_heading(doc, "Vérifier", 2)
    add_bullets(doc, [
        "Ai je chuté ou subi un choc important",
        "Ai je une faiblesse nouvelle, une perte de sensibilité au périnée ou un trouble urinaire",
        "Ai je de la fièvre, un malaise important ou un symptôme inhabituel",
    ], checkbox=True)
    add_heading(doc, "S’installer", 2)
    add_body(doc, "Testez deux ou trois positions sans chercher la posture parfaite. Par exemple, assis avec les pieds soutenus, debout appuyé sur un plan stable ou allongé avec les jambes soutenues. Gardez la position qui vous permet de relâcher l’effort pendant quelques minutes.")
    add_heading(doc, "Préparer la suite", 2)
    add_bullets(doc, [
        "placez le téléphone et l’eau à portée",
        "dégagez le trajet vers les toilettes",
        "demandez de l’aide pour les charges ou les soins aux enfants si nécessaire",
        "notez l’heure d’apparition et les symptômes associés",
    ])

    new_page(doc, "Jour 1", "Protéger sans immobiliser", "Alterner repos bref et activité tolérable")
    add_body(doc, "Le repos au lit prolongé n’est généralement pas recommandé pour une lombalgie commune. Cela ne signifie pas qu’il faut ignorer la douleur. Alternez quelques minutes de repos confortable avec des déplacements courts et utiles.")
    add_heading(doc, "Le rythme du jour", 2)
    add_bullets(doc, [
        "matin  se lever en plusieurs temps et faire un premier déplacement court",
        "milieu de journée  changer de position avant une raideur prolongée",
        "après midi  répéter un trajet simple plutôt que faire une longue sortie",
        "soir  préparer la nuit et noter les changements",
    ])
    add_heading(doc, "Ce que vous observez", 2)
    add_lines(doc, ["Position la plus confortable", "Déplacement le plus facile", "Symptôme nouveau ou changement", "Décision pour demain"], 1)

    new_page(doc, "Jour 1", "Se lever se coucher et s’habiller", "Utiliser des appuis temporaires sans en faire une règle rigide")
    add_heading(doc, "Sortir du lit", 2)
    add_body(doc, "Essayez de rouler sur le côté, de rapprocher les jambes du bord puis de pousser avec les bras pour vous asseoir. Si une autre façon est plus confortable et sûre, utilisez la. Le but est de répartir l’effort, pas de respecter une technique parfaite.")
    add_heading(doc, "S’habiller", 2)
    add_body(doc, "Asseyez vous pour enfiler pantalon, chaussettes et chaussures. Rapprochez le vêtement plutôt que de chercher une grande amplitude. Une pince de préhension ou l’aide d’un proche peut être utile pendant quelques jours.")
    add_heading(doc, "Utiliser une béquille ou une canne", 2)
    add_body(doc, "Une aide à la marche mal réglée peut augmenter l’effort ou le risque de chute. Si vous en avez déjà une, utilisez la selon les consignes reçues. Si vous pensez en avoir besoin pour la première fois, demandez conseil à un professionnel plutôt que d’improviser le réglage.")
    add_image(doc, "bequilles-mobilite.png", "Aide à la marche  Le réglage et l’apprentissage doivent être adaptés à la personne", width=12.8)

    new_page(doc, "Jour 2", "Retrouver un peu de mobilité", "Choisir des mouvements connus et de faible amplitude")
    add_body(doc, "Commencez par les mouvements nécessaires à votre journée. Se retourner, se lever, marcher jusqu’à une autre pièce et changer d’appui sont déjà des activités. Il n’est pas indispensable d’ajouter une longue séance.")
    add_heading(doc, "Routine prototype", 2)
    add_numbered(doc, [
        "Respirez calmement dans une position confortable pendant quelques cycles.",
        "Effectuez un petit transfert de poids ou un mouvement du bassin qui reste contrôlable.",
        "Marchez une à trois minutes dans un espace sûr.",
        "Reposez vous brièvement puis notez la réponse.",
    ])
    add_note(doc, "Validation", "La version finale comportera des illustrations de trois options et leurs contre indications. Cette routine ne doit pas être publiée sans validation clinique.")
    add_lines(doc, ["Option choisie", "Durée", "Réponse immédiate", "Réponse plus tard"], 1)

    new_page(doc, "Jour 3", "Reprendre une tâche utile", "Mesurer la capacité par une action réelle")
    add_body(doc, "Choisissez une tâche courte qui vous rend un peu d’autonomie, comme préparer un repas simple, prendre une douche ou ranger un petit objet. Réduisez la tâche avant de commencer et prévoyez une pause.")
    add_heading(doc, "Fractionner la tâche", 2)
    add_numbered(doc, [
        "Préparez le matériel à hauteur accessible.",
        "Effectuez une première partie pendant quelques minutes.",
        "Changez de position ou reposez vous brièvement.",
        "Décidez si vous terminez, déléguez ou reportez la suite.",
    ])
    add_lines(doc, ["Tâche choisie", "Partie réalisée", "Aide utilisée", "Ce que je modifierai"], 1)

    new_page(doc, "Jour 4", "Sortir et se déplacer", "Préparer le trajet avant de tester la distance")
    add_heading(doc, "Avant de partir", 2)
    add_bullets(doc, [
        "choisir un trajet court avec un point de retour",
        "prévoir un siège ou un lieu de pause",
        "éviter de porter une charge non nécessaire",
        "informer quelqu’un si vous manquez de confiance",
    ], checkbox=True)
    add_heading(doc, "En voiture", 2)
    add_body(doc, "Rapprochez le siège pour éviter d’aller chercher les pédales, gardez les objets utiles accessibles et interrompez les longs trajets si possible. Pour entrer, asseyez vous d’abord puis faites pivoter les jambes ensemble si cette option vous convient.")
    add_heading(doc, "Dans les transports", 2)
    add_body(doc, "Privilégiez un horaire moins chargé, utilisez les appuis et acceptez de vous asseoir. La priorité est la sécurité dans un environnement imprévisible.")

    new_page(doc, "Jour 5", "Préparer la reprise du travail", "Décrire les contraintes avant de choisir une date")
    add_body(doc, "La reprise dépend de votre état, du trajet, des horaires et des tâches réelles. Une activité professionnelle peut parfois être poursuivie ou reprise avec des adaptations temporaires. Discutez en avec votre médecin et, si nécessaire, avec le médecin du travail.")
    add_heading(doc, "Inventaire concret", 2)
    add_lines(doc, [
        "Temps de trajet et mode de transport",
        "Durée des positions prolongées",
        "Charges ou gestes répétitifs",
        "Possibilités d’alterner les tâches",
        "Adaptation temporaire à demander",
    ], 1)
    add_note(doc, "Conseil", "Ne transmettez pas de détail médical à votre employeur au delà de ce qui est nécessaire. Le médecin du travail peut proposer des adaptations en respectant le secret médical.")

    new_page(doc, "Jour 6", "Préparer une consultation", "Rendre visibles les changements et les limites")
    add_image(doc, "consultation-kine.png", "Un échange utile relie les symptômes aux activités et aux décisions", width=13.8)
    add_heading(doc, "À apporter", 2)
    add_bullets(doc, [
        "liste des médicaments et traitements déjà essayés",
        "heure ou date de début et évolution sur les derniers jours",
        "description des symptômes dans la jambe s’il y en a",
        "activités impossibles, difficiles et encore possibles",
        "vos deux questions prioritaires",
    ], checkbox=True)
    add_lines(doc, ["Question 1", "Question 2", "Ce que je veux retrouver en priorité"], 1)

    new_page(doc, "Jour 7", "Faire le bilan et choisir la suite", "Amélioration stagnation ou aggravation")
    add_heading(doc, "Si la situation s’améliore", 2)
    add_body(doc, "Conservez les activités tolérées et augmentez un seul paramètre à la fois. Le cahier Dos solide 28 jours peut servir de cadre après validation de votre situation.")
    add_heading(doc, "Si la situation stagne", 2)
    add_body(doc, "Regardez les capacités autant que la douleur. Si les activités restent très limitées ou si vous êtes inquiet, prenez un avis médical plutôt que de multiplier les exercices au hasard.")
    add_heading(doc, "Si la situation s’aggrave", 2)
    add_body(doc, "Revenez à la page de sécurité. Un symptôme nouveau, une faiblesse, une atteinte du périnée, un trouble urinaire, une fièvre ou une altération générale impose une évaluation adaptée.")
    add_lines(doc, ["Ce qui a progressé", "Ce qui reste difficile", "Ma décision pour les sept prochains jours"], 1)

    new_page(doc, "Vie quotidienne", "Quand les contraintes s’accumulent", "Douleur organisation et pression financière")
    add_image(doc, "factures-pression.png", "Une poussée douloureuse peut compliquer les démarches le travail et le budget", width=13.8)
    add_body(doc, "La douleur peut arriver au milieu de responsabilités qui ne s’arrêtent pas. Faites d’abord une liste courte des conséquences des sept prochains jours, puis séparez ce qui doit être fait, délégué ou reporté.")
    add_heading(doc, "Démarches utiles", 2)
    add_bullets(doc, [
        "conserver les ordonnances, arrêts, factures et justificatifs au même endroit",
        "demander à l’organisme concerné les conditions de prise en charge avant une dépense importante",
        "prévenir rapidement si un rendez vous ou un paiement ne peut pas être tenu",
        "solliciter un service social si la santé met le logement, le revenu ou l’accès aux soins en difficulté",
    ])
    add_note(doc, "Limite", "Ce guide ne fournit pas de conseil juridique, assurantiel ou financier personnalisé. Les droits et remboursements dépendent de la situation et du régime applicable.")

    new_page(doc, "Références", "Sources et validation éditoriale")
    add_sources(doc)
    add_heading(doc, "Étapes de production restantes", 2)
    add_bullets(doc, [
        "relecture médicale du parcours des sept jours",
        "création de schémas pour les transferts et les positions de repos",
        "enregistrement de sept audios courts à partir du texte validé",
        "test utilisateur en situation de douleur avec critères d’arrêt stricts",
    ])
    out = OUT / "03-SOS-lumbago-7-jours-manuscrit-v0.1.docx"
    doc.save(out)
    return out


def sheet_start(doc, number, title, subtitle):
    new_page(doc, f"Fiche {number:02d}", title, subtitle)


def build_sheets():
    doc = base_document("Fiches pratiques")
    cover(
        doc,
        "Collection détachable",
        "Fiches pratiques Le Réflexe Dos",
        "Lot 01 pour les premiers jours la consultation et la reprise",
        "consultation-kine.png",
        "Une collection conçue pour être imprimée ou consultée séparément",
        "Huit fiches indépendantes à garder sur le téléphone, imprimer ou transmettre à un proche. Chaque fiche tient sur une situation précise.",
    )
    add_safety_page(doc)

    new_page(doc, "Mode d’emploi", "Une fiche pour une décision")
    add_body(doc, "Commencez par la fiche 02 si la douleur est récente, inhabituelle ou plus forte que d’habitude. Ensuite, choisissez seulement la fiche qui répond à votre besoin du moment. Vous n’avez pas à tout lire ni à tout faire.")
    add_note(doc, "Notre ton", "Ces fiches sont là pour vous simplifier la prochaine décision. Elles reconnaissent vos contraintes, proposent plusieurs options et vous encouragent à demander de l’aide lorsque vous en avez besoin.")
    add_heading(doc, "Le gabarit commun", 2)
    add_bullets(doc, [
        "objectif de la fiche",
        "contrôle de sécurité",
        "étapes courtes",
        "adaptation possible",
        "règle d’arrêt",
        "note personnelle",
    ])
    add_heading(doc, "Les huit fiches", 2)
    add_numbered(doc, [
        "Que faire dans les trente premières minutes",
        "Quand demander de l’aide",
        "Préparer une consultation",
        "Organiser une journée difficile",
        "Bouger quand on travaille assis",
        "Préparer la nuit",
        "Se déplacer avec une aide à la marche",
        "Faire le bilan de la semaine",
    ])

    sheet_start(doc, 1, "Les trente premières minutes", "Réduire l’incertitude avant d’agir")
    add_heading(doc, "Objectif", 2)
    add_body(doc, "Vérifier la sécurité, trouver une position de répit et organiser le prochain déplacement utile.")
    add_heading(doc, "Étapes", 2)
    add_numbered(doc, [
        "Repérez l’heure et les circonstances d’apparition.",
        "Vérifiez les signes de la fiche 02.",
        "Testez deux positions sûres pendant quelques minutes.",
        "Gardez à portée le téléphone, l’eau et les objets indispensables.",
        "Faites un déplacement court si cela reste possible et sûr.",
    ])
    add_heading(doc, "Adaptation", 2)
    add_body(doc, "Demandez à un proche de sécuriser le trajet, d’apporter les objets ou de prendre en charge une tâche urgente. Évitez de porter une charge pour prouver que vous pouvez encore le faire.")
    add_heading(doc, "Règle d’arrêt", 2)
    add_body(doc, "Un trouble urinaire, une perte de sensibilité du périnée, une faiblesse progressive, un traumatisme important, une fièvre ou une altération générale impose une évaluation urgente ou rapide.")
    add_lines(doc, ["Ma position de répit", "La prochaine action utile"], 1)

    sheet_start(doc, 2, "Quand demander de l’aide", "Une décision de sécurité avant toute routine")
    add_heading(doc, "Aide urgente", 2)
    add_bullets(doc, [
        "difficulté nouvelle à uriner ou perte du contrôle des urines ou des selles",
        "perte de sensibilité autour du périnée",
        "faiblesse importante ou qui progresse dans une ou deux jambes",
        "douleur après un traumatisme important",
        "altération marquée de l’état général ou symptômes inhabituels inquiétants",
    ], checkbox=True)
    add_body(doc, "En France, contactez le 15 ou le 112 si la situation paraît urgente ou si vous ne savez pas quel service appeler.")
    add_heading(doc, "Avis médical rapide", 2)
    add_bullets(doc, [
        "fièvre, douleur nocturne incessante ou aggravation progressive",
        "perte de poids inexpliquée, antécédent de cancer ou traitement prolongé par corticoïdes",
        "engourdissement ou faiblesse nouvelle dans une jambe",
        "forte limitation des activités sans amélioration",
    ], checkbox=True)
    add_note(doc, "À savoir", "Cette liste aide à décider de demander un avis. Elle ne permet ni d’identifier ni d’exclure une maladie.")

    sheet_start(doc, 3, "Préparer une consultation", "Donner au professionnel les informations qui changent la décision")
    add_heading(doc, "Avant le rendez vous", 2)
    add_bullets(doc, [
        "notez la date de début et les circonstances",
        "décrivez la zone et une éventuelle douleur dans la jambe",
        "listez les symptômes associés et leurs changements",
        "apportez les traitements, examens et antécédents utiles",
        "choisissez deux questions prioritaires",
    ], checkbox=True)
    add_lines(doc, ["Ce que je ne peux plus faire", "Ce que je peux encore faire", "Question 1", "Question 2"], 1)
    add_heading(doc, "Après le rendez vous", 2)
    add_lines(doc, ["Ce que j’ai compris", "Ce que je fais aujourd’hui", "Quand je dois reconsulter"], 1)

    sheet_start(doc, 4, "Organiser une journée difficile", "Préserver l’essentiel sans rester immobile toute la journée")
    add_heading(doc, "Choisir trois priorités", 2)
    add_lines(doc, ["Indispensable aujourd’hui", "Peut être délégué", "Peut attendre"], 1)
    add_heading(doc, "Créer un rythme", 2)
    add_bullets(doc, [
        "alterner une tâche courte et un repos bref",
        "changer de position avant une longue période figée",
        "garder un déplacement simple dans la journée",
        "préparer les objets avant de commencer une tâche",
    ], checkbox=True)
    add_heading(doc, "Réduire la charge mentale", 2)
    add_body(doc, "Regroupez les justificatifs, ordonnances et factures dans une même pochette ou un dossier numérique. Demandez une information écrite avant une dépense importante si la prise en charge est incertaine.")
    add_lines(doc, ["La personne que je peux appeler", "Le réglage le plus utile aujourd’hui"], 1)

    sheet_start(doc, 5, "Bouger quand on travaille assis", "Varier les positions plutôt que chercher la posture parfaite")
    add_heading(doc, "Objectif", 2)
    add_body(doc, "Réduire les longues périodes sans changement et conserver une activité professionnelle compatible avec la situation.")
    add_heading(doc, "Plan de la demi journée", 2)
    add_bullets(doc, [
        "pieds soutenus et écran lisible sans tension inutile",
        "rappels de changement de position",
        "appels réalisés debout si cela est plus confortable",
        "courte marche ou déplacement entre deux blocs de travail",
    ], checkbox=True)
    add_heading(doc, "Adaptations à discuter", 2)
    add_bullets(doc, [
        "alternance temporaire des tâches",
        "réduction des charges ou des déplacements",
        "horaires ou reprise progressive selon avis médical",
        "évaluation par le médecin du travail si la reprise est difficile",
    ])
    add_lines(doc, ["Mon changement de position le plus simple", "L’adaptation à demander"], 1)

    sheet_start(doc, 6, "Préparer la nuit", "Chercher le confort sans imposer une position unique")
    add_heading(doc, "Avant de vous coucher", 2)
    add_bullets(doc, [
        "préparez le trajet nocturne et un éclairage accessible",
        "placez les objets utiles à portée sans encombrer le passage",
        "testez un coussin sous ou entre les jambes si cela améliore le confort",
        "préparez une façon de vous relever en plusieurs temps",
    ], checkbox=True)
    add_heading(doc, "Si le sommeil est interrompu", 2)
    add_body(doc, "Changez de position, levez vous quelques instants si cela vous aide, puis revenez au lit. Évitez d’interpréter une mauvaise nuit comme la preuve d’une aggravation. Demandez un avis si la douleur nocturne est incessante, s’aggrave ou s’accompagne d’autres signes d’alerte.")
    add_lines(doc, ["Position la plus confortable", "Ce qui facilite le lever"], 1)

    sheet_start(doc, 7, "Utiliser une aide à la marche", "Sécuriser le déplacement avant de chercher la distance")
    add_image(doc, "bequilles-mobilite.png", "Une aide utile doit être réglée et utilisée correctement", width=12.2)
    add_heading(doc, "Avant l’utilisation", 2)
    add_bullets(doc, [
        "vérifiez l’état des embouts et la stabilité",
        "portez des chaussures stables",
        "dégagez les tapis et obstacles",
        "utilisez l’aide selon le réglage et les consignes reçues",
    ], checkbox=True)
    add_body(doc, "Une canne ou des béquilles utilisées pour la première fois méritent un conseil professionnel. La hauteur, le côté d’utilisation et la séquence de marche dépendent de votre situation.")
    add_heading(doc, "Règle d’arrêt", 2)
    add_body(doc, "Arrêtez si l’aide augmente le risque de chute, provoque un malaise ou s’accompagne d’une faiblesse nouvelle. Demandez de l’aide pour les escaliers si vous n’avez pas appris la technique adaptée.")

    sheet_start(doc, 8, "Faire le bilan de la semaine", "Transformer les observations en une seule décision")
    add_lines(doc, [
        "Une activité devenue plus facile",
        "Une activité encore évitée",
        "Une réponse du corps devenue plus prévisible",
        "Un réglage qui m’aide",
        "Un signe qui justifie un avis",
    ], 1)
    add_heading(doc, "Décision", 2)
    add_bullets(doc, [
        "Je conserve la même dose une semaine de plus.",
        "J’augmente un seul paramètre.",
        "Je réduis et je demande conseil.",
        "Je consulte car les symptômes ont changé ou restent très limitants.",
    ], checkbox=True)
    add_lines(doc, ["Mon prochain rendez vous avec moi même", "L’action choisie"], 1)

    new_page(doc, "Références", "Sources et validation éditoriale")
    add_sources(doc)
    add_heading(doc, "Étapes de production restantes", 2)
    add_bullets(doc, [
        "faire valider chaque règle d’arrêt par un professionnel de santé",
        "créer une version une page recto pour les fiches 1 à 6",
        "tester l’impression en noir et blanc et sur téléphone",
        "ajouter un QR vers les mises à jour et la date de révision",
    ])
    out = OUT / "04-Fiches-pratiques-lot-01-manuscrit-v0.1.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    outputs = [build_ebook(), build_workbook(), build_sos(), build_sheets()]
    for path in outputs:
        print(path)
